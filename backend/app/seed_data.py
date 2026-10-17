"""Load the CS study plan and the users workbook."""
import json
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from .config import BASE_DIR
from .models import DEAN, FACULTY, HEAD, TEAM, Course, Department, Term, User

# Department names as they appear in the college faculty list.
DEPARTMENTS = {
    "علوم الحاسبات": "CS",
    "النظم والشبكات": "IS",
    "تقنية المعلومات": "IT",
    "هندسة البرمجيات": "SE",
}
ROLE_FROM_SHEET = {"عميد الكلية": DEAN, "رئيس قسم": HEAD, "عضو هيئة تدريس": FACULTY, "فريق المشروع": TEAM}


def seed_departments(db: Session) -> None:
    for name, program in DEPARTMENTS.items():
        if not db.scalar(select(Department).where(Department.name == name)):
            db.add(Department(name=name, program_code=program))
    db.flush()


def seed_plan(db: Session, path: Path = BASE_DIR / "seed" / "cs_plan.json") -> int:
    plan = json.loads(path.read_text(encoding="utf-8"))
    added = 0
    for order, c in enumerate(plan["courses"]):
        exists = db.scalar(select(Course).where(Course.program == plan["program"], Course.title == c["title"]))
        if exists:
            continue
        db.add(Course(program=plan["program"], code=c.get("code"), spec_code=c.get("spec_code"), title=c["title"],
                      level=c.get("level"), kind=c["kind"], group_label=c.get("group"), credit_hours=c.get("ch"),
                      practical=c.get("practical", "unknown"), sort_order=order))
        added += 1
    db.flush()
    return added


def seed_term(db: Session, code: str = "1447-1") -> Term:
    term = db.scalar(select(Term).where(Term.code == code))
    if not term:
        for t in db.scalars(select(Term)):
            t.active = False
        term = Term(code=code, active=True)
        db.add(term)
        db.flush()
    return term


def import_users_workbook(db: Session, source) -> dict:
    """Sheet "المستخدمون": name, email, email status, rank, department, role, in pilot, courses."""
    seed_departments(db)
    depts = {d.name: d for d in db.scalars(select(Department))}
    ws = load_workbook(source, read_only=True)["المستخدمون"]
    added = updated = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        name = (row[0] or "").strip()
        if not name or name.startswith("("):
            continue
        email, email_status, rank, dept_name, role_ar, pilot = row[1], row[2] or "", row[3], (row[4] or "").strip(), row[5], row[6]
        user = db.scalar(select(User).where(User.name == name))
        if not user:
            user = User(name=name)
            db.add(user)
            added += 1
        else:
            updated += 1
        user.email = (email or "").strip() or None
        user.email_confirmed = email_status == "مؤكد"
        user.rank = rank
        user.department = depts.get(dept_name)
        user.role = ROLE_FROM_SHEET.get(role_ar or "", FACULTY)
        user.in_pilot = pilot == "نعم"
    db.flush()
    return {"added": added, "updated": updated}
