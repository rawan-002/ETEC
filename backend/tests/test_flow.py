import io
import zipfile

from app.models import Course, Document, Notification, User
from tests.conftest import login


def docx_bytes(text: str) -> bytes:
    import docx
    d = docx.Document()
    for line in text.split("\n"):
        d.add_paragraph(line)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def pptx_bytes(text: str) -> bytes:
    from pptx import Presentation
    p = Presentation()
    s = p.slides.add_slide(p.slide_layouts[1])
    s.shapes.title.text = text
    buf = io.BytesIO()
    p.save(buf)
    return buf.getvalue()


def course(db, code) -> Course:
    return db.query(Course).filter_by(code=code).one()


def select(client, items):
    r = client.put("/api/me/selection", json={"items": items})
    assert r.status_code == 200, r.text
    return r.json()


# ---------- login ----------

def test_login_list_and_faculty_login_without_pin(client, db):
    users = client.get("/api/auth/users").json()
    assert users[0]["role"] == "dean" and users[0]["needs_pin"] is True
    assert not next(u for u in users if u["name"] == "أستاذ أ")["needs_pin"]
    login(client, db, "أستاذ أ")
    me = client.get("/api/auth/me").json()
    assert me["name"] == "أستاذ أ" and me["needs_course_selection"] is True


def test_dean_pin_required_and_lockout(client, db):
    dean = db.query(User).filter_by(name="العميد").one()
    assert client.post("/api/auth/login", json={"user_id": dean.id}).status_code == 401
    for _ in range(3):
        assert client.post("/api/auth/login", json={"user_id": dean.id, "pin": "000000"}).status_code == 401
    assert client.post("/api/auth/login", json={"user_id": dean.id, "pin": "000000"}).status_code == 423
    # Locked even with the right PIN until the lock expires.
    assert client.post("/api/auth/login", json={"user_id": dean.id, "pin": "123456"}).status_code == 423


def test_dean_correct_pin(client, db):
    login(client, db, "العميد", "123456")
    assert client.get("/api/auth/me").json()["role"] == "dean"


# ---------- course selection ----------

def test_selection_and_coordinator_conflict(client, db):
    c = course(db, "CS1258")
    login(client, db, "أستاذ أ")
    assert select(client, [{"course_id": c.id, "is_coordinator": True}])["conflicts"] == []
    login(client, db, "أستاذ ب")
    res = select(client, [{"course_id": c.id, "is_coordinator": True}])
    assert res["conflicts"][0]["coordinator"] == "أستاذ أ"
    teachers = next(x for x in client.get("/api/courses").json() if x["id"] == c.id)["teachers"]
    assert {(t["name"], t["is_coordinator"]) for t in teachers} == {("أستاذ أ", True), ("أستاذ ب", False)}


def test_empty_selection_counts_as_answered_and_head_is_not_forced(client, db):
    login(client, db, "أستاذ أ")
    assert client.get("/api/auth/me").json()["needs_course_selection"] is True
    select(client, [])
    assert client.get("/api/auth/me").json()["needs_course_selection"] is False
    login(client, db, "رئيس CS")
    assert client.get("/api/auth/me").json()["needs_course_selection"] is False


def test_register_request_needs_head_approval(client, db):
    from app.models import Department
    cs = db.query(Department).filter_by(program_code="CS").one()
    assert client.post("/api/auth/register-request", json={"name": "د. عضو جديد للتجربة", "department_id": cs.id}).status_code == 200
    assert "د. عضو جديد للتجربة" not in [u["name"] for u in client.get("/api/auth/users").json()]
    login(client, db, "رئيس CS")
    pending = client.get("/api/programs/CS/distribution").json()["pending_users"]
    assert pending[0]["name"] == "د. عضو جديد للتجربة"
    assert client.post(f"/api/users/{pending[0]['id']}/approve").status_code == 200
    assert "د. عضو جديد للتجربة" in [u["name"] for u in client.get("/api/auth/users").json()]


# ---------- upload & classification ----------

def upload(client, course_id, name, data, component=None):
    files = [("files", (name, data))]
    r = client.post(f"/api/courses/{course_id}/upload", files=files, data={"component": component} if component else {})
    assert r.status_code == 200, r.text
    return r.json()["documents"]


def test_upload_classifies_by_name_and_content(client, db):
    c = course(db, "CS1258")
    login(client, db, "أستاذ أ")
    select(client, [{"course_id": c.id, "is_coordinator": True}])
    spec = docx_bytes("Course Specifications\nCourse Code: CS1005\nCourse Learning Outcomes\nCredit Hours: 3")
    upload(client, c.id, "document1.docx", spec)
    upload(client, c.id, "Midterm.docx", docx_bytes("Midterm exam\nStudent name:\nAnswer all questions"))

    detail = client.get(f"/api/courses/{c.id}").json()
    status = {x["component"]: x["status"] for x in detail["components"]}
    assert status["spec"] == "present" and status["assessment"] == "present"
    assert status["theory"] == "missing"
    # Specification code differs from the plan code: recorded as a finding.
    assert any("CS1005" in f["message"] for f in detail["findings"])


def test_zip_upload_skips_temp_files_and_flags_broken_ones(client, db):
    c = course(db, "CS1261")
    login(client, db, "أستاذ أ")
    select(client, [{"course_id": c.id, "is_coordinator": True}])
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("Lectures/Lecture 1.pptx", pptx_bytes("Lecture 1 - Introduction"))
        z.writestr("Labs/Lab_5.pdf", b"\x00" * 100)
        z.writestr("Labs/~$Lab_6.docx", b"temp")
    ids = upload(client, c.id, "course.zip", buf.getvalue())
    assert len(ids) == 2
    db.expire_all()
    docs = {d.filename: d for d in db.query(Document).filter(Document.id.in_(ids))}
    assert docs["Lecture 1.pptx"].component == "theory" and docs["Lecture 1.pptx"].status == "classified"
    assert docs["Lab_5.pdf"].status == "error"


def test_word_file_named_pdf_is_detected(client, db):
    c = course(db, "CS1775")
    login(client, db, "أستاذ أ")
    select(client, [{"course_id": c.id, "is_coordinator": True}])
    upload(client, c.id, "TP-153 Course Specifications-CS1509.pdf", docx_bytes("Course Specifications\nCredit Hours"))
    detail = client.get(f"/api/courses/{c.id}").json()
    assert any(f["kind"] == "wrong_extension" for f in detail["findings"])


def test_duplicate_upload_is_marked(client, db):
    c = course(db, "CS1004")
    login(client, db, "أستاذ أ")
    select(client, [{"course_id": c.id, "is_coordinator": True}])
    data = docx_bytes("Final exam\nAnswer all questions")
    upload(client, c.id, "final.docx", data)
    upload(client, c.id, "final copy.docx", data)
    db.expire_all()
    assert [d.status for d in db.query(Document).filter_by(course_id=c.id).order_by(Document.id)] == ["classified", "duplicate"]


def test_low_confidence_goes_to_team_and_team_resolves(client, db):
    c = course(db, "CS1004")
    login(client, db, "أستاذ أ")
    select(client, [{"course_id": c.id, "is_coordinator": True}])
    (doc_id,) = upload(client, c.id, "file1.docx", docx_bytes("some notes"))
    login(client, db, "فريق")
    queue = client.get("/api/review").json()
    assert queue[0]["id"] == doc_id
    assert client.post(f"/api/documents/{doc_id}/review", json={"component": "theory"}).status_code == 200
    assert client.get("/api/review").json() == []


# ---------- permissions ----------

def test_head_scope_and_faculty_scope(client, db):
    c = course(db, "CS1258")
    login(client, db, "رئيس SE")
    assert client.get("/api/programs/CS/distribution").status_code == 403
    assert client.post("/api/programs/CS/distribution/approve").status_code == 403
    assert client.get(f"/api/courses/{c.id}").status_code == 403
    login(client, db, "أستاذ ب")
    assert client.get(f"/api/courses/{c.id}").status_code == 403
    assert client.post(f"/api/courses/{c.id}/upload", files=[("files", ("a.pdf", b"%PDF-1.4"))]).status_code == 403
    assert client.get("/api/users").status_code == 403
    login(client, db, "رئيس CS")
    assert client.get(f"/api/courses/{c.id}").status_code == 200


# ---------- reminders ----------

def test_reminders_start_after_approval_and_stop_when_complete(client, db):
    c = course(db, "CS1004")  # practical N/A: spec, theory, assessment needed
    login(client, db, "أستاذ أ")
    select(client, [{"course_id": c.id, "is_coordinator": True}])
    login(client, db, "العميد", "123456")
    assert client.post("/api/reminders/run").json()["sent"] == 0  # distribution not approved yet

    login(client, db, "رئيس CS")
    client.post("/api/programs/CS/distribution/approve")
    login(client, db, "العميد", "123456")
    assert client.post("/api/reminders/run").json()["sent"] == 1
    n = db.query(Notification).filter_by(kind="missing_reminder").one()
    assert "CS1004" in n.body and n.status == "simulated"

    login(client, db, "أستاذ أ")
    upload(client, c.id, "spec.docx", docx_bytes("Course Specifications\nCourse Learning Outcomes"), "spec")
    upload(client, c.id, "Lecture 1.pptx", pptx_bytes("Lecture 1"), "theory")
    upload(client, c.id, "final exam.docx", docx_bytes("Final exam"), "assessment")
    assert client.get(f"/api/courses/{c.id}").json()["percent"] == 100
    login(client, db, "العميد", "123456")
    assert client.post("/api/reminders/run").json()["sent"] == 0


def test_dashboard_and_activity_scoping(client, db):
    login(client, db, "العميد", "123456")
    dash = client.get("/api/dashboard").json()
    cs = next(p for p in dash["programs"] if p["program"] == "CS")
    assert cs["courses_total"] == 59
    login(client, db, "أستاذ ب")
    assert client.get("/api/dashboard").json()["programs"] == []
    msgs = [a["message"] for a in client.get("/api/activity").json()]
    assert all("العميد" not in m for m in msgs)


# ---------- sandbox testers ----------

def test_tester_hidden_login_and_full_isolation(client, db):
    db.add(User(name="مختبر", role="faculty", is_tester=True, in_pilot=True, login_alias="taif"))
    db.commit()
    c = course(db, "CS1004")
    # A real member uploads first; the tester must never see this file.
    login(client, db, "أستاذ أ")
    select(client, [{"course_id": c.id, "is_coordinator": True}])
    (real_doc,) = upload(client, c.id, "spec.docx", docx_bytes("Course Specifications\nCourse Learning Outcomes"))

    client.cookies.clear()
    assert "مختبر" not in [u["name"] for u in client.get("/api/auth/users").json()]
    tester_id = db.query(User).filter_by(name="مختبر").one().id
    assert client.post("/api/auth/login", json={"user_id": tester_id}).status_code == 404
    assert client.post("/api/auth/alias", json={"word": "nobody"}).status_code == 404
    assert client.post("/api/auth/alias", json={"word": " Taif "}).status_code == 200
    assert client.get("/api/auth/me").json()["is_tester"] is True

    select(client, [{"course_id": c.id, "is_coordinator": True}])
    upload(client, c.id, "Final exam.docx", docx_bytes("Final exam\nAnswer all questions"))
    mine = client.get(f"/api/courses/{c.id}").json()
    assert [d["filename"] for d in mine["documents"]] == ["Final exam.docx"]
    status = {x["component"]: x["status"] for x in mine["components"]}
    assert status["assessment"] == "present" and status["spec"] == "missing"
    assert client.get(f"/api/documents/{real_doc}/download").status_code == 404
    assert all("أستاذ أ" not in a["message"] for a in client.get("/api/activity").json())

    # Real users see none of the tester's data.
    login(client, db, "أستاذ أ")
    real = client.get(f"/api/courses/{c.id}").json()
    assert [d["filename"] for d in real["documents"]] == ["spec.docx"]
    assert [t["name"] for t in real["teachers"]] == ["أستاذ أ"]
    assert all("مختبر" not in a["message"] for a in client.get("/api/activity").json())
    login(client, db, "رئيس CS")
    client.post("/api/programs/CS/distribution/approve")
    dist = client.get("/api/programs/CS/distribution").json()
    assert all(t["name"] != "مختبر" for x in dist["courses"] for t in x["teachers"])
    login(client, db, "العميد", "123456")
    client.post("/api/reminders/run")
    assert not db.query(Notification).join(User).filter(User.name == "مختبر").count()
